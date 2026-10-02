from pathlib import Path
import re
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(tmp_path: Path):
    try:
        from main import create_app
    except (ImportError, ModuleNotFoundError) as exc:
        pytest.fail(f"application factory is not implemented: {exc}")
    app = create_app(f"sqlite:///{(tmp_path / 'web-test.db').as_posix()}")
    with TestClient(app) as test_client:
        yield test_client


@pytest.mark.parametrize("path", ["/", "/components", "/bom-library", "/storage", "/boms", "/orders", "/inventory", "/data-management"])
def test_main_pages_render(client, path):
    response = client.get(path)

    assert response.status_code == 200
    assert "元件管理" in response.text


def test_component_api_creates_manual_component(client):
    response = client.post(
        "/api/components",
        json={"comment": "10kΩ ±1%", "footprint": "0603", "supplier_part": "C1001"},
    )

    assert response.status_code == 201
    assert response.json()["comment"] == "10kΩ ±1%"
    assert response.json()["supplier"] == "LCSC"


def test_component_api_round_trips_category_and_description(client):
    created = client.post(
        "/api/components",
        json={
            "supplier_part": "C8734",
            "comment": "STM32F103C8T6",
            "footprint": "LQFP-48",
            "secondary_category": "单片机(MCU/MPU/SOC)",
            "description": "CPU内核:ARM Cortex-M3;CPU最大主频:72MHz;",
        },
    )
    assert created.status_code == 201
    assert created.json()["secondary_category"] == "单片机(MCU/MPU/SOC)"
    assert "72MHz" in created.json()["description"]

    updated = client.patch(
        f"/api/components/{created.json()['id']}",
        json={"description": "更新后的说明"},
    )
    assert updated.status_code == 200
    assert updated.json()["description"] == "更新后的说明"


def test_storage_api_creates_named_custom_grid(client):
    response = client.post("/api/storage/boxes", json={"name": "24 格黑盒", "rows": 4, "columns": 6})

    assert response.status_code == 201
    assert response.json()["name"] == "24 格黑盒"
    assert response.json()["cell_count"] == 24


def test_bom_library_preview_does_not_write_and_confirm_only_creates_library(client, bom_file):
    with bom_file.open("rb") as handle:
        preview = client.post(
            "/api/bom-library/preview",
            files={"file": ("sample_bom.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert preview.status_code == 200
    assert preview.json()["items"][0]["supplier_part"] == "C1001"
    assert client.get("/api/components").json() == []

    confirmed = client.post(
        "/api/bom-library/confirm",
        json={"token": preview.json()["token"]},
    )

    assert confirmed.status_code == 201
    assert confirmed.json()["created_count"] == 2
    assert confirmed.json()["reused_count"] == 0
    components = client.get("/api/components").json()
    assert {item["supplier_part"] for item in components} == {"C1001", "C1002"}
    assert all(item["available_quantity"] == 0 for item in components)
    assert "控制板" not in client.get("/boms").text


def test_bom_library_preview_exposes_field_conflicts_and_confirm_applies_selection(client, tmp_path):
    client.post(
        "/api/components",
        json={
            "supplier_part": "C1001",
            "comment": "10kΩ",
            "footprint": "R0603",
            "secondary_category": "旧分类",
            "description": "旧说明",
        },
    )
    path = tmp_path / "metadata.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "Supplier Part", "Comment", "Footprint", "Quantity",
        "Secondary Category", "Description",
    ])
    sheet.append(["C1001", "10kΩ", "R0603", 1, "贴片电阻", "新说明"])
    workbook.save(path)
    with path.open("rb") as handle:
        preview = client.post(
            "/api/bom-library/preview",
            files={
                "file": (
                    path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )

    assert preview.status_code == 200
    assert {
        (value["field"], value["current_value"], value["incoming_value"])
        for value in preview.json()["conflicts"]
    } == {
        ("secondary_category", "旧分类", "贴片电阻"),
        ("description", "旧说明", "新说明"),
    }

    confirmed = client.post(
        "/api/bom-library/confirm",
        json={
            "token": preview.json()["token"],
            "overwrites": [{"supplier_part": "C1001", "fields": ["description"]}],
        },
    )
    assert confirmed.status_code == 201
    component = client.get("/api/components").json()[0]
    assert component["secondary_category"] == "旧分类"
    assert component["description"] == "新说明"


def test_bom_library_conflict_controls_default_to_checked(client):
    page = client.get("/components?show_bom_import=1").text

    assert "bom-library-conflicts" in page
    assert "data-select-all-conflicts" in page
    assert "data-clear-all-conflicts" in page


def test_bom_library_script_defaults_conflict_checkboxes_to_checked():
    script = Path("static/js/bom_library.js").read_text(encoding="utf-8")
    assert "data-overwrite-conflict" in script
    assert "checked" in script
    assert "data-select-all-conflicts" in script
    assert "data-clear-all-conflicts" in script
    assert "input.checked = checked" in script


def test_component_page_shows_searchable_metadata_in_expandable_detail(client):
    created = client.post(
        "/api/components",
        json={
            "supplier_part": "C8734",
            "comment": "STM32F103C8T6",
            "footprint": "LQFP-48",
            "secondary_category": "单片机(MCU/MPU/SOC)",
            "description": "CPU内核:ARM Cortex-M3;CPU最大主频:72MHz;CPU位数:32 Bit;",
        },
    ).json()

    page = client.get("/components").text

    assert "单片机(MCU/MPU/SOC)" in page
    assert "CPU最大主频:72MHz" in page
    assert 'class="metadata-description"' in page
    assert f'data-component-row="{created["id"]}"' in page
    assert 'name="secondary_category"' in page
    assert 'name="description"' in page


def test_component_page_paginates_fifty_records(client):
    for index in range(51):
        response = client.post(
            "/api/components",
            json={
                "supplier_part": f"P{index:03d}",
                "comment": "分页测试",
                "footprint": "0603",
            },
        )
        assert response.status_code == 201

    first = client.get("/components")
    second = client.get("/components?page=2")
    below_range = client.get("/components?page=0")
    above_range = client.get("/components?page=999")

    assert first.text.count('data-component-group="') == 50
    assert "共 51 种" in first.text
    assert 'href="/components?q=&amp;page=2"' in first.text
    assert 'data-component-group="51"' not in first.text
    assert second.text.count('data-component-group="') == 1
    assert 'data-component-group="51"' in second.text
    assert re.search(r'aria-current="page"[^>]*>1</a>', below_range.text)
    assert re.search(r'aria-current="page"[^>]*>2</a>', above_range.text)


def test_component_page_searches_before_paginating(client):
    for index in range(51):
        response = client.post(
            "/api/components",
            json={
                "supplier_part": f"MATCH-{index:03d}",
                "comment": "高压 电阻",
                "footprint": "0603",
            },
        )
        assert response.status_code == 201
    client.post(
        "/api/components",
        json={"supplier_part": "OTHER", "comment": "电容", "footprint": "0805"},
    )

    first = client.get("/components?q=高压")
    second = client.get("/components?q=高压&page=2")

    assert first.text.count('data-component-group="') == 50
    assert "共 51 种" in first.text
    assert "OTHER" not in first.text
    assert "q=%E9%AB%98%E5%8E%8B&amp;page=2" in first.text
    assert second.text.count('data-component-group="') == 1


def test_component_page_server_search_includes_metadata_and_location(client):
    target = client.post(
        "/api/components",
        json={
            "supplier_part": "TARGET",
            "comment": "精密元件",
            "footprint": "QFN-32",
            "manufacturer_part": "MFG-UNIQUE",
            "secondary_category": "特殊分类",
            "description": "耐高温详细说明",
        },
    ).json()
    client.post(
        "/api/components",
        json={"supplier_part": "OTHER", "comment": "普通元件", "footprint": "0603"},
    )
    box = client.post(
        "/api/storage/boxes",
        json={"name": "目标收纳盒", "rows": 1, "columns": 1},
    ).json()
    assigned = client.post(
        f"/api/storage/cells/{box['cells'][0]['id']}/assign-component",
        json={"component_id": target["id"]},
    )
    assert assigned.status_code == 200

    for keyword in ["MFG-UNIQUE", "特殊分类", "耐高温", "目标收纳盒", "A1"]:
        page = client.get("/components", params={"q": keyword})
        assert "TARGET" in page.text
        assert "OTHER" not in page.text
        assert "共 1 种" in page.text


def test_bom_page_shows_snapshot_and_candidate_metadata(client, tmp_path):
    client.post(
        "/api/components",
        json={
            "supplier_part": "C1001",
            "comment": "10kΩ",
            "footprint": "R0603",
            "secondary_category": "库存精确料分类",
            "description": "库存精确料说明",
        },
    )
    candidate = client.post(
        "/api/components",
        json={
            "supplier_part": "C9001",
            "comment": "10kΩ",
            "footprint": "ALT-0603",
            "secondary_category": "候选分类",
            "description": "候选库存说明",
        },
    ).json()
    client.post(
        f"/api/components/{candidate['id']}/adjust-stock",
        json={"delta": 10, "note": "候选库存"},
    )

    path = tmp_path / "pick-metadata.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "Supplier Part", "Comment", "Footprint", "Quantity", "Designator",
        "Secondary Category", "Description",
    ])
    sheet.append([
        "C1001", "10kΩ", "R0603", 2, "R1,R2", "BOM 分类", "BOM 原始说明",
    ])
    workbook.save(path)
    with path.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={
                "file": (
                    path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "资料展示测试"},
    ).json()

    page = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text
    assert "BOM 分类" in page
    assert "BOM 原始说明" in page
    assert "候选分类" in page
    assert "候选库存说明" in page
    assert "待替代物料资料" in page
    assert "与待替代物料不同" in page


def test_bom_metadata_uses_regular_label_when_exact_stock_is_enough(client, tmp_path):
    component = client.post(
        "/api/components",
        json={
            "supplier_part": "C1001",
            "comment": "10kΩ",
            "footprint": "R0603",
            "secondary_category": "贴片电阻",
            "description": "阻值:10kΩ;精度:±1%;",
        },
    ).json()
    client.post(
        f"/api/components/{component['id']}/adjust-stock",
        json={"delta": 10, "note": "充足库存"},
    )

    path = tmp_path / "enough-stock.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "Supplier Part", "Comment", "Footprint", "Quantity",
        "Secondary Category", "Description",
    ])
    sheet.append(["C1001", "10kΩ", "R0603", 2, "贴片电阻", "阻值:10kΩ;精度:±1%;"])
    workbook.save(path)
    with path.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={
                "file": (
                    path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "库存充足资料展示"},
    ).json()

    page = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text
    assert "物料资料" in page
    assert "待替代物料资料" not in page
    assert "贴片电阻" in page
    assert "阻值:10kΩ;精度:±1%;" in page
    assert "参数匹配候选" not in page


def test_old_bom_uses_original_component_metadata_after_substitution(client, tmp_path):
    client.post("/api/components", json={
        "supplier_part": "C1001",
        "comment": "100nF",
        "footprint": "C0805",
        "secondary_category": "原料分类",
        "description": "原料 Description",
    })
    candidate = client.post("/api/components", json={
        "supplier_part": "C9001",
        "comment": "100nF",
        "footprint": "C0603",
        "secondary_category": "候选分类",
        "description": "候选 Description",
    }).json()
    client.post(
        f"/api/components/{candidate['id']}/adjust-stock",
        json={"delta": 10, "note": "候选库存"},
    )
    path = tmp_path / "legacy.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Supplier Part", "Comment", "Footprint", "Quantity", "Designator"])
    sheet.append(["C1001", "100nF", "C0805", 2, "C1,C2"])
    workbook.save(path)
    with path.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={
                "file": (
                    path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "旧 BOM"},
    ).json()
    item = bom["items"][0]

    before = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text
    assert "原料分类" in before
    assert "原料 Description" in before
    assert "元件库补显" in before
    assert "待替代物料资料" in before

    selected = client.post(
        f"/api/boms/items/{item['id']}/substitute",
        json={"component_id": candidate["id"]},
    )
    assert selected.status_code == 200
    after = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text
    assert "原料分类" in after
    assert "原料 Description" in after
    assert "候选分类" in after
    assert "候选 Description" in after
    assert "待替代物料资料" in after


def test_bom_snapshot_metadata_takes_priority_over_library_fallback(client, tmp_path):
    client.post("/api/components", json={
        "supplier_part": "C1001",
        "comment": "100nF",
        "footprint": "C0805",
        "secondary_category": "库内分类",
        "description": "库内 Description",
    })
    path = tmp_path / "snapshot.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "Supplier Part", "Comment", "Footprint", "Quantity",
        "Secondary Category", "Description",
    ])
    sheet.append(["C1001", "100nF", "C0805", 1, "BOM 分类", "BOM Description"])
    workbook.save(path)
    with path.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={
                "file": (
                    path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "快照优先"},
    ).json()

    page = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text
    assert "BOM 分类" in page
    assert "BOM Description" in page
    assert "库内分类" not in page
    assert "库内 Description" not in page


def test_component_and_bom_pages_group_rows_by_component(client, tmp_path):
    client.post("/api/components", json={
        "supplier_part": "C1001",
        "comment": "100nF",
        "footprint": "C0805",
        "secondary_category": "贴片电容",
        "description": "100nF 50V X7R",
    })
    component_page = client.get("/components").text
    assert '<tbody class="component-record"' in component_page

    path = tmp_path / "grouped.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Supplier Part", "Comment", "Footprint", "Quantity"])
    sheet.append(["C1001", "100nF", "C0805", 1])
    workbook.save(path)
    with path.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={
                "file": (
                    path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "分组测试"},
    ).json()
    bom_page = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text
    assert '<tbody class="bom-item-record"' in bom_page
    assert "待替代物料资料" in bom_page
    styles = Path("static/css/style.css").read_text(encoding="utf-8")
    assert ".component-record" in styles
    assert ".bom-item-record" in styles
    assert ".bom-source-metadata" in styles
    assert ".component-record:nth-of-type(even)" in styles
    assert ".bom-item-record:nth-of-type(even)" in styles
    assert "--record-main-bg" in styles
    assert "--record-metadata-bg" in styles
    assert "--record-main-bg: #e7f0ff" in styles
    assert "--record-metadata-bg: #d8e6fa" in styles
    assert "--record-metadata-bg: #f3f4f6" in styles
    assert "--record-marker: #c6cbd3" in styles
    assert "--record-marker: #79a6ed" in styles
    assert "--record-border: #d8dde4" in styles
    assert ".candidate-panel" in styles


def test_pick_bom_import_does_not_create_unknown_components(client, bom_file):
    with bom_file.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={"file": ("sample_bom.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    confirmed = client.post(
        "/api/boms/confirm",
        json={"token": preview.json()["token"], "name": "控制板"},
    )

    assert confirmed.status_code == 201
    assert confirmed.json()["item_count"] == 2
    assert client.get("/api/components").json() == []
    assert "未建档" in client.get(f"/boms?bom_id={confirmed.json()['id']}").text


def test_library_preview_token_cannot_confirm_pick_bom(client, bom_file):
    with bom_file.open("rb") as handle:
        preview = client.post(
            "/api/bom-library/preview",
            files={"file": ("sample_bom.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    response = client.post(
        "/api/boms/confirm",
        json={"token": preview.json()["token"], "name": "不应导入"},
    )

    assert response.status_code == 409
    assert client.get("/api/components").json() == []


def _create_stocked_pick_bom(client, bom_file):
    required = {"C1001": 5, "C1002": 1}
    components = []
    for supplier_part, quantity in required.items():
        component = client.post(
            "/api/components",
            json={"comment": supplier_part, "footprint": "0603", "supplier_part": supplier_part},
        ).json()
        components.append(component)
        client.post(
            f"/api/components/{component['id']}/adjust-stock",
            json={"delta": quantity + 2, "note": "测试入库"},
        )
    box = client.post(
        "/api/storage/boxes",
        json={"name": "领料测试盒", "rows": 1, "columns": 2},
    ).json()
    for component, cell in zip(components, box["cells"], strict=True):
        client.post(
            f"/api/storage/cells/{cell['id']}/assign-component",
            json={"component_id": component["id"]},
        )
    with bom_file.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={"file": ("sample_bom.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "逐格领料测试"},
    ).json()
    return bom, required


def test_pick_bom_item_endpoint_deducts_only_selected_item(client, bom_file):
    bom, required = _create_stocked_pick_bom(client, bom_file)
    first_item = next(item for item in bom["items"] if item["supplier_part"] == "C1001")

    response = client.post(f"/api/boms/items/{first_item['id']}/pick")

    assert response.status_code == 200
    assert response.json()["picked"] is True
    assert response.json()["bom_status"] == "open"
    quantities = {item["supplier_part"]: item["available_quantity"] for item in client.get("/api/components").json()}
    assert quantities["C1001"] == 2
    assert quantities["C1002"] == required["C1002"] + 2


def test_pick_bom_item_endpoint_rejects_repeat(client, bom_file):
    bom, _ = _create_stocked_pick_bom(client, bom_file)
    item_id = bom["items"][0]["id"]
    assert client.post(f"/api/boms/items/{item_id}/pick").status_code == 200

    repeated = client.post(f"/api/boms/items/{item_id}/pick")

    assert repeated.status_code == 409
    assert "已经领取" in repeated.json()["detail"]


def test_bom_item_can_confirm_comment_matched_substitute(client, bom_file):
    candidate = client.post(
        "/api/components",
        json={
            "comment": "  10KΩ  ",
            "footprint": "DIFFERENT-FOOTPRINT",
            "supplier_part": "C9001",
        },
    ).json()
    client.post(
        f"/api/components/{candidate['id']}/adjust-stock",
        json={"delta": 20, "note": "候选库存"},
    )
    box = client.post(
        "/api/storage/boxes",
        json={"name": "候选盒", "rows": 1, "columns": 1},
    ).json()
    client.post(
        f"/api/storage/cells/{box['cells'][0]['id']}/assign-component",
        json={"component_id": candidate["id"]},
    )
    with bom_file.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={"file": ("sample_bom.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "替代接口测试"},
    ).json()
    item = next(value for value in bom["items"] if value["supplier_part"] == "C1001")

    response = client.post(
        f"/api/boms/items/{item['id']}/substitute",
        json={"component_id": candidate["id"]},
    )

    assert response.status_code == 200
    assert response.json()["component_id"] == candidate["id"]
    assert response.json()["original_supplier_part"] == "C1001"
    assert response.json()["selected_supplier_part"] == "C9001"
    updated = next(value for value in client.get("/api/components").json() if value["id"] == candidate["id"])
    assert updated["available_quantity"] == 20


def test_bom_page_shows_comment_candidates_and_disables_unusable_choices(client, bom_file):
    def create_candidate(part, footprint, quantity):
        component = client.post(
            "/api/components",
            json={"comment": "10KΩ", "footprint": footprint, "supplier_part": part},
        ).json()
        if quantity:
            client.post(
                f"/api/components/{component['id']}/adjust-stock",
                json={"delta": quantity, "note": "候选库存"},
            )
        return component

    enough = create_candidate("C9001", "ALT-ENOUGH", 20)
    short = create_candidate("C9002", "ALT-SHORT", 1)
    unlocated = create_candidate("C9003", "ALT-UNLOCATED", 20)
    box = client.post(
        "/api/storage/boxes",
        json={"name": "候选盒", "rows": 1, "columns": 2},
    ).json()
    for component, cell in zip([enough, short], box["cells"], strict=True):
        client.post(
            f"/api/storage/cells/{cell['id']}/assign-component",
            json={"component_id": component["id"]},
        )
    with bom_file.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={"file": ("sample_bom.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "候选界面测试"},
    ).json()
    item = next(value for value in bom["items"] if value["supplier_part"] == "C1001")

    page = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text

    assert "参数匹配候选" in page
    assert "Footprint 不参与匹配，请人工确认" in page
    assert "ALT-ENOUGH" in page
    assert "ALT-SHORT" in page
    assert "ALT-UNLOCATED" in page
    assert f'data-substitute-item="{item["id"]}"' in page
    assert f'data-component-id="{enough["id"]}"' in page
    assert f'data-component-id="{short["id"]}"' not in page
    assert f'data-component-id="{unlocated["id"]}"' in page

    confirmed = client.post(
        f"/api/boms/items/{item['id']}/substitute",
        json={"component_id": unlocated["id"]},
    )
    assert confirmed.status_code == 200
    confirmed_page = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text
    assert "实际领料：C9003" in confirmed_page
    assert "已选择替代 · 未定位" in confirmed_page
    assert 'class="card box-card"' not in confirmed_page


def test_bom_page_hides_comment_candidates_when_exact_stock_is_enough(client, bom_file):
    exact = client.post(
        "/api/components",
        json={"comment": "10kΩ", "footprint": "R0603", "supplier_part": "C1001"},
    ).json()
    client.post(
        f"/api/components/{exact['id']}/adjust-stock",
        json={"delta": 5, "note": "精确库存"},
    )
    client.post(
        "/api/components",
        json={"comment": "10kΩ", "footprint": "ALT", "supplier_part": "C9001"},
    )
    with bom_file.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={"file": ("sample_bom.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "精确料测试"},
    ).json()

    page = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text

    assert "参数匹配候选" not in page


def test_substitute_selection_and_highlight_are_isolated_between_boms(client, bom_file):
    candidate = client.post(
        "/api/components",
        json={"comment": "10KΩ", "footprint": "ALT", "supplier_part": "C9001"},
    ).json()
    client.post(
        f"/api/components/{candidate['id']}/adjust-stock",
        json={"delta": 20, "note": "候选库存"},
    )
    box = client.post(
        "/api/storage/boxes",
        json={"name": "隔离测试盒", "rows": 1, "columns": 1},
    ).json()
    client.post(
        f"/api/storage/cells/{box['cells'][0]['id']}/assign-component",
        json={"component_id": candidate["id"]},
    )

    def import_pick_bom(name):
        with bom_file.open("rb") as handle:
            preview = client.post(
                "/api/boms/preview",
                files={"file": ("sample_bom.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            ).json()
        return client.post(
            "/api/boms/confirm",
            json={"token": preview["token"], "name": name},
        ).json()

    first = import_pick_bom("表 1")
    second = import_pick_bom("表 2")
    first_item = next(value for value in first["items"] if value["supplier_part"] == "C1001")
    second_item = next(value for value in second["items"] if value["supplier_part"] == "C1001")
    selected = client.post(
        f"/api/boms/items/{first_item['id']}/substitute",
        json={"component_id": candidate["id"]},
    )
    assert selected.status_code == 200

    first_page = client.get(f"/boms?bom_id={first['id']}&expanded=1").text
    second_page = client.get(f"/boms?bom_id={second['id']}&expanded=1").text

    assert "实际领料：C9001" in first_page
    assert f'data-pick-item="{first_item["id"]}"' in first_page
    assert "实际领料：C9001" not in second_page
    assert f'data-substitute-item="{second_item["id"]}"' in second_page
    assert f'data-pick-item="{second_item["id"]}"' not in second_page
    assert 'class="card box-card"' not in second_page


def test_substitute_and_native_demand_share_one_pick_cell(client, tmp_path):
    path = tmp_path / "shared-component-bom.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Supplier Part", "Comment", "Footprint", "Quantity", "Designator"])
    sheet.append(["A", "10kΩ", "R0603", 3, "R1,R2,R3"])
    sheet.append(["B", "10kΩ", "R0805", 5, "R4,R5,R6,R7,R8"])
    workbook.save(path)

    client.post(
        "/api/components",
        json={"supplier_part": "A", "comment": "10kΩ", "footprint": "R0603"},
    )
    candidate = client.post(
        "/api/components",
        json={"supplier_part": "B", "comment": "10kΩ", "footprint": "R0805"},
    ).json()
    client.post(
        f"/api/components/{candidate['id']}/adjust-stock",
        json={"delta": 10, "note": "聚合领料测试"},
    )
    box = client.post(
        "/api/storage/boxes",
        json={"name": "共享格位盒", "rows": 1, "columns": 1},
    ).json()
    client.post(
        f"/api/storage/cells/{box['cells'][0]['id']}/assign-component",
        json={"component_id": candidate["id"]},
    )
    with path.open("rb") as handle:
        preview = client.post(
            "/api/boms/preview",
            files={
                "file": (
                    path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        ).json()
    bom = client.post(
        "/api/boms/confirm",
        json={"token": preview["token"], "name": "共享格位 BOM"},
    ).json()
    source_item = next(value for value in bom["items"] if value["supplier_part"] == "A")
    native_item = next(value for value in bom["items"] if value["supplier_part"] == "B")
    selected = client.post(
        f"/api/boms/items/{source_item['id']}/substitute",
        json={"component_id": candidate["id"]},
    )
    assert selected.status_code == 200

    page = client.get(f"/boms?bom_id={bom['id']}&expanded=1").text

    assert page.count('class="storage-cell highlight pickable-cell"') == 1
    assert "共取 8 件" in page
    assert "A → B" in page
    assert "R1, R2, R3" in page
    assert "<b>B</b>：5 件" in page
    assert "R4, R5, R6, R7, R8" in page

    picked = client.post(f"/api/boms/items/{source_item['id']}/pick")
    assert picked.status_code == 200
    assert picked.json()["bom_status"] == "picked"
    components = {
        value["supplier_part"]: value
        for value in client.get("/api/components").json()
    }
    assert components["A"]["available_quantity"] == 0
    assert components["B"]["available_quantity"] == 2
    assert source_item["id"] != native_item["id"]


def test_highlighted_cell_content_uses_selected_bom_item_component():
    from app.routes.pages import templates
    from app.services.bom import BomPickGroup

    cell_component = SimpleNamespace(supplier_part="CELL-A", comment="cell component")
    cell = SimpleNamespace(id=10, label="A1", component=cell_component)
    selected_component = SimpleNamespace(
        id=2,
        supplier_part="BOM-B",
        comment="selected bom component",
        available_quantity=10,
        storage_cell=cell,
    )
    item = SimpleNamespace(
        id=20,
        component_id=2,
        component=selected_component,
        supplier_part="ORIGINAL-B",
        comment="selected bom component",
        footprint="0603",
        quantity=2,
        warning="",
        designators="R1, R2",
    )
    bom = SimpleNamespace(
        id=30,
        name="表 2",
        file_name="bom.xlsx",
        status="open",
        items=[item],
    )
    box = SimpleNamespace(id=40, name="测试盒", columns=1, cells=[cell])
    cell.box = box
    group = BomPickGroup(
        component=selected_component,
        items=[item],
        picked_item_ids=set(),
    )

    html = templates.env.get_template("boms.html").render(
        request=SimpleNamespace(url=SimpleNamespace(path="/boms")),
        boms=[bom],
        selected=bom,
        boxes=[box],
        pick_groups={cell.id: group},
        pending_quantity_by_component={selected_component.id: item.quantity},
        picked_item_ids=set(),
        candidates_by_item={},
        candidate_requirements={},
        expanded=True,
    )
    grid_html = html.split('class="storage-grid"', 1)[1]

    assert "BOM-B" in grid_html
    assert "selected bom component" in grid_html
    assert "CELL-A" not in grid_html
    assert "cell component" not in grid_html


def test_legacy_whole_bom_pick_endpoint_is_removed(client, bom_file):
    bom, _ = _create_stocked_pick_bom(client, bom_file)

    response = client.post(f"/api/boms/{bom['id']}/pick")

    assert response.status_code == 404


def test_pick_interface_has_collapsed_details_and_cell_confirmation(client, bom_file):
    bom, _ = _create_stocked_pick_bom(client, bom_file)
    first_item = bom["items"][0]

    page = client.get(f"/boms?bom_id={bom['id']}&collapsed=1").text

    assert re.search(r'<details class="card bom-detail"(?![^>]*\bopen\b)[^>]*>', page)
    assert f'data-pick-item="{first_item["id"]}"' in page
    assert 'class="pick-overlay"' in page
    assert "确认领料" in page
    assert "取消" in page
    assert "确认领料并扣库" not in page


def test_picked_cell_renders_permanent_overlay_without_action(client, bom_file):
    bom, _ = _create_stocked_pick_bom(client, bom_file)
    first_item = bom["items"][0]
    client.post(f"/api/boms/items/{first_item['id']}/pick")

    page = client.get(f"/boms?bom_id={bom['id']}").text

    assert f'data-picked-item="{first_item["id"]}"' in page
    assert "已领取" in page
    assert f'data-pick-item="{first_item["id"]}"' not in page


def test_order_pending_item_can_be_placed_and_stocked(client, order_file):
    component = client.post(
        "/api/components",
        json={"comment": "10k", "footprint": "0603", "supplier_part": "C1001"},
    ).json()
    box = client.post("/api/storage/boxes", json={"name": "电阻盒", "rows": 1, "columns": 1}).json()

    with order_file.open("rb") as handle:
        preview = client.post(
            "/api/orders/preview",
            files={"file": ("sample_order.xls", handle, "application/vnd.ms-excel")},
        )
    confirmed = client.post("/api/orders/confirm", json={"token": preview.json()["token"]})

    assert confirmed.status_code == 201
    pending = next(item for item in confirmed.json()["items"] if item["supplier_part"] == "C1001")
    unmatched = next(item for item in confirmed.json()["items"] if item["supplier_part"] == "C9999")
    assert pending["status"] == "pending_location"
    assert unmatched["status"] == "unmatched"

    placed = client.post(
        f"/api/orders/items/{pending['id']}/place",
        json={"cell_id": box["cells"][0]["id"]},
    )
    assert placed.status_code == 200
    assert placed.json()["status"] == "stocked"
    updated = next(item for item in client.get("/api/components").json() if item["id"] == component["id"])
    assert updated["available_quantity"] == 8


def test_unmatched_order_item_can_be_linked_to_manual_component(client, order_file):
    manual = client.post(
        "/api/components",
        json={"comment": "1uF", "footprint": "0603"},
    ).json()
    with order_file.open("rb") as handle:
        preview = client.post(
            "/api/orders/preview",
            files={"file": ("sample_order.xls", handle, "application/vnd.ms-excel")},
        )
    order = client.post("/api/orders/confirm", json={"token": preview.json()["token"]}).json()
    item = next(value for value in order["items"] if value["supplier_part"] == "C9999")

    linked = client.post(
        f"/api/orders/items/{item['id']}/link",
        json={"component_id": manual["id"]},
    )

    assert linked.status_code == 200
    assert linked.json()["status"] == "pending_location"


def test_duplicate_order_confirmation_is_rejected(client, order_file):
    def preview_order():
        with order_file.open("rb") as handle:
            return client.post(
                "/api/orders/preview",
                files={"file": ("sample_order.xls", handle, "application/vnd.ms-excel")},
            ).json()

    first = preview_order()
    assert client.post("/api/orders/confirm", json={"token": first["token"]}).status_code == 201
    second = preview_order()
    duplicate = client.post("/api/orders/confirm", json={"token": second["token"]})

    assert duplicate.status_code == 409
    assert "已经导入" in duplicate.json()["detail"]


def test_inventory_page_can_reverse_manual_adjustment(client):
    component = client.post(
        "/api/components",
        json={"comment": "10k", "footprint": "0603", "supplier_part": "C1001"},
    ).json()
    client.post(
        f"/api/components/{component['id']}/adjust-stock",
        json={"delta": 9, "note": "首次盘点"},
    )

    page = client.get("/inventory?q=C1001")

    assert page.status_code == 200
    assert "库存流水" in page.text
    assert "首次盘点" in page.text
    movement_id = int(re.search(r'data-reverse-movement="(\d+)"', page.text).group(1))

    reversed_response = client.post(f"/api/inventory/movements/{movement_id}/reverse")

    assert reversed_response.status_code == 200
    assert reversed_response.json()["quantity_delta"] == -9
    updated = next(item for item in client.get("/api/components").json() if item["id"] == component["id"])
    assert updated["available_quantity"] == 0
    assert client.post(f"/api/inventory/movements/{movement_id}/reverse").status_code == 409


def test_data_management_downloads_excel_and_database_backup(client):
    client.post(
        "/api/components",
        json={"comment": "10k", "footprint": "0603", "supplier_part": "C1001"},
    )

    excel = client.get("/api/data/export.xlsx")
    backup = client.get("/api/data/backup")

    assert excel.status_code == 200
    assert excel.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert excel.content.startswith(b"PK")
    assert backup.status_code == 200
    assert backup.content.startswith(b"SQLite format 3")


def test_data_management_restore_replaces_database_snapshot(client):
    client.post(
        "/api/components",
        json={"comment": "旧快照", "footprint": "0603", "supplier_part": "OLD"},
    )
    backup = client.get("/api/data/backup").content
    client.post(
        "/api/components",
        json={"comment": "稍后新增", "footprint": "0603", "supplier_part": "NEW"},
    )

    restored = client.post(
        "/api/data/restore",
        files={"file": ("snapshot.db", backup, "application/octet-stream")},
    )

    assert restored.status_code == 200
    assert "恢复" in restored.json()["message"]
    assert {item["supplier_part"] for item in client.get("/api/components").json()} == {"OLD"}


def test_data_management_restore_rejects_invalid_database(client):
    response = client.post(
        "/api/data/restore",
        files={"file": ("broken.db", b"not sqlite", "application/octet-stream")},
    )

    assert response.status_code == 409
    assert "SQLite" in response.json()["detail"]
