from io import BytesIO

import openpyxl
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.db import Base
from app.models import Bom, BomItem, Component, StockMovement, StorageBox, StorageCell
from app.services import BusinessRuleError


def _create_file_database(path, component_part: str):
    engine = create_engine(f"sqlite:///{path.as_posix()}")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(Component(
            supplier="LCSC",
            supplier_part=component_part,
            comment="10k",
            footprint="0603",
        ))
        session.commit()
    engine.dispose()


def test_build_excel_export_contains_five_business_sheets(db_session):
    from app.services.data_management import build_excel_export

    component = Component(
        supplier="LCSC",
        supplier_part="C1001",
        comment="10k",
        footprint="0603",
        secondary_category="贴片电阻",
        description="厚膜电阻 10kΩ ±1%",
        available_quantity=8,
    )
    box = StorageBox(name="电阻盒", rows=1, columns=1)
    box.cells.append(StorageCell(row_index=0, column_index=0, label="A1", component=component))
    db_session.add_all([component, box])
    db_session.flush()
    db_session.add(StockMovement(component=component, quantity_delta=8, movement_type="adjustment"))
    bom = Bom(name="控制板", file_name="control.xlsx")
    bom.items.append(BomItem(
        supplier_part="C1001",
        comment="10k",
        footprint="0603",
        secondary_category="BOM 电阻分类",
        description="BOM 原始参数说明",
        quantity=2,
        designators="R1, R2",
    ))
    db_session.add(bom)
    db_session.commit()

    workbook = openpyxl.load_workbook(BytesIO(build_excel_export(db_session)))

    assert workbook.sheetnames == ["元件库", "收纳盒", "库存流水", "订单", "BOM"]
    assert workbook["元件库"]["B2"].value == "C1001"
    assert workbook["收纳盒"]["A2"].value == "电阻盒"
    assert workbook["库存流水"]["D2"].value == 8
    component_headers = [cell.value for cell in workbook["元件库"][1]]
    assert "二级分类" in component_headers
    assert "Description" in component_headers
    category_column = component_headers.index("二级分类") + 1
    description_column = component_headers.index("Description") + 1
    assert workbook["元件库"].cell(2, category_column).value == "贴片电阻"
    assert "厚膜电阻" in workbook["元件库"].cell(2, description_column).value

    bom_headers = [cell.value for cell in workbook["BOM"][1]]
    bom_category_column = bom_headers.index("二级分类") + 1
    bom_description_column = bom_headers.index("Description") + 1
    assert workbook["BOM"].cell(2, bom_category_column).value == "BOM 电阻分类"
    assert workbook["BOM"].cell(2, bom_description_column).value == "BOM 原始参数说明"


def test_create_database_backup_returns_readable_sqlite_snapshot(tmp_path):
    from app.services.data_management import create_database_backup, validate_database

    database_path = tmp_path / "current.db"
    _create_file_database(database_path, "C1001")

    backup = create_database_backup(f"sqlite:///{database_path.as_posix()}")
    snapshot = tmp_path / "snapshot.db"
    snapshot.write_bytes(backup)

    validate_database(snapshot)
    assert backup.startswith(b"SQLite format 3")


def test_restore_database_replaces_data_and_creates_safety_backup(tmp_path):
    from app.services.data_management import restore_database

    current_path = tmp_path / "current.db"
    uploaded_path = tmp_path / "uploaded.db"
    backup_directory = tmp_path / "backups"
    _create_file_database(current_path, "OLD")
    _create_file_database(uploaded_path, "NEW")

    safety_backup = restore_database(
        f"sqlite:///{current_path.as_posix()}",
        uploaded_path,
        backup_directory,
    )

    engine = create_engine(f"sqlite:///{current_path.as_posix()}")
    with Session(engine) as session:
        parts = session.scalars(select(Component.supplier_part)).all()
    engine.dispose()
    assert parts == ["NEW"]
    assert safety_backup.parent == backup_directory
    assert safety_backup.exists()
    assert safety_backup.read_bytes().startswith(b"SQLite format 3")


def test_restore_database_rejects_invalid_file_without_changing_current_data(tmp_path):
    from app.services.data_management import restore_database

    current_path = tmp_path / "current.db"
    uploaded_path = tmp_path / "invalid.db"
    _create_file_database(current_path, "C1001")
    uploaded_path.write_text("not sqlite", encoding="utf-8")

    with pytest.raises(BusinessRuleError, match="有效的 SQLite"):
        restore_database(
            f"sqlite:///{current_path.as_posix()}",
            uploaded_path,
            tmp_path / "backups",
        )

    engine = create_engine(f"sqlite:///{current_path.as_posix()}")
    with Session(engine) as session:
        assert session.scalar(select(Component.supplier_part)) == "C1001"
    engine.dispose()
