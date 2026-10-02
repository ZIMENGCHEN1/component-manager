from pathlib import Path

import openpyxl
import pytest


def require_bom_importer():
    try:
        from app.importers.bom import ImportValidationError, parse_bom
    except ModuleNotFoundError as exc:
        pytest.fail(f"BOM importer is not implemented: {exc}")
    return parse_bom, ImportValidationError


def test_bom_parser_merges_duplicate_supplier_parts(bom_file):
    parse_bom, _ = require_bom_importer()
    result = parse_bom(bom_file)

    assert len(result.items) == 2
    first = result.items[0]
    assert first.supplier_part == "C1001"
    assert first.quantity == 5
    assert first.designators == ["R1", "R2", "R3", "R4", "R5"]


def test_bom_parser_reports_missing_columns(tmp_path: Path):
    parse_bom, ImportValidationError = require_bom_importer()
    path = tmp_path / "bad.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Comment", "Quantity"])
    workbook.save(path)

    with pytest.raises(ImportValidationError, match="Supplier Part"):
        parse_bom(path)


def test_bom_parser_reads_optional_category_and_description(tmp_path: Path):
    parse_bom, _ = require_bom_importer()
    path = tmp_path / "metadata.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "Supplier Part", "Comment", "Footprint", "Quantity",
        "Secondary Category", "Description",
    ])
    sheet.append(["C8734", "STM32F103C8T6", "LQFP-48", 1, "单片机", "72MHz Cortex-M3"])
    workbook.save(path)

    item = parse_bom(path).items[0]

    assert item.secondary_category == "单片机"
    assert item.description == "72MHz Cortex-M3"


def test_bom_parser_keeps_optional_metadata_empty_for_legacy_file(bom_file):
    parse_bom, _ = require_bom_importer()

    item = parse_bom(bom_file).items[0]

    assert item.secondary_category == ""
    assert item.description == ""


def test_duplicate_part_fills_blank_metadata_and_warns_on_nonblank_difference(tmp_path: Path):
    parse_bom, _ = require_bom_importer()
    path = tmp_path / "duplicates.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "Supplier Part", "Comment", "Footprint", "Quantity", "Designator",
        "Secondary Category", "Description",
    ])
    sheet.append(["C1", "10k", "0603", 1, "R1", "", "厚膜电阻"])
    sheet.append(["C1", "10k", "0603", 1, "R2", "贴片电阻", "薄膜电阻"])
    workbook.save(path)

    result = parse_bom(path)
    item = result.items[0]

    assert item.secondary_category == "贴片电阻"
    assert item.description == "厚膜电阻"
    assert any("Description" in value for value in item.warnings)
    assert any("Description" in value for value in result.warnings)
